"""Reconcile the KR experiment with cached values in the two original files.

Only anonymous counts and contrasts are printed. Does not edit workbooks.
Column choices below are intentionally independent of experiment_data.BLOCKS.
"""
import json
import sqlite3
from pathlib import Path
from statistics import mean
from openpyxl import load_workbook
from experiment_data import build_experiment

root = Path(__file__).resolve().parent.parent
con = sqlite3.connect(root / "Оценки_Арабули_2012-2026.sqlite")
con.row_factory = sqlite3.Row
experiment = build_experiment(con)
raw = {}
for group, size in (("9Б", 26), ("9Д", 28)):
    wb = load_workbook(root / f"{group}.xlsx", data_only=True, read_only=True)
    rows = list(wb["Уроки"].iter_rows(min_row=2, max_row=size+1, values_only=True))
    class_scores = {}
    for row_index, row in enumerate(rows):
        scores = {}
        for block, points_col, grade_cols in ((2,21,(22,23)), (3,30,(31,32)),
                                              (4,57,(58,59)), (5,68,(69,70)), (6,77,(78,79))):
            points = row[points_col-1]
            marks = [row[col-1] for col in grade_cols]
            if isinstance(points,(int,float)) and points >= 0 and all(mark in (2,3,4,5) for mark in marks):
                scores[block] = mean(marks)
        class_scores[row_index] = scores
    raw[group] = class_scores
    wb.close()

checks = []
kr = next(metric for metric in experiment["metrics"] if metric["id"] == "kr")
for sample_id in ("balanced", "available"):
    selected = {group: [scores for scores in students.values()
                       if sample_id == "available" or len(scores) == 5]
                for group, students in raw.items()}
    sample = next(item for item in kr["samples"] if item["id"] == sample_id)
    gaps = {}
    for block in range(2,7):
        result_row = next(row for row in sample["timeline"] if row["block"] == block)
        values = {group: [s[block] for s in scores if block in s] for group,scores in selected.items()}
        for group in ("9Б", "9Д"):
            assert len(values[group]) == result_row["groups"][group]["n"]
            assert abs(mean(values[group])-result_row["groups"][group]["mean"]) < 1e-6
        gaps[block] = mean(values["9Б"]) - mean(values["9Д"])
        assert abs(gaps[block]-result_row["gap"]) < 1e-6
    pooled = (gaps[5]-gaps[4])/2
    assert abs(pooled-sample["effects"][2]["value"]) < 1e-6
    checks.append({"sample":sample_id,"pupils":{g:len(v) for g,v in selected.items()},"pooled":round(pooled,6)})
for risk in experiment["risk"]:
    assert risk["lower"] + risk["middle"] + risk["higher"] == risk["n"]
    effect = kr["samples"][0]["effects"][0 if risk["group"] == "9Д" else 1]["value"]
    assert abs(mean(risk["changes"])-effect) < 2e-6
print(json.dumps({"sourceReconciliation":"passed", "checks":checks}, ensure_ascii=False))
